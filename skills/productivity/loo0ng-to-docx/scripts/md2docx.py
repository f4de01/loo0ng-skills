#!/usr/bin/env python3
"""冻结转换器：把模型写的 Markdown 最小集转成 DOCX，以该节点的官方模板为版式载体。只依赖 python-docx，离线。

依据 ADR-0006。做法：打开模板、清 body 只留 sectPr、按模板里现成段落的段落格式与字体写段，表格照抄模板的
tblPr、列宽、行高与单元格格式；合并单元格由 Markdown 约定表达（`<` 与左格合并、`^` 与上格合并）。收尾两条确定性
规则（通用裁定台账 #10、#11）：清元数据（作者、最后修改者置空，修订号置 1，删上次打印时间，创建与修改时间置为本次）；
正文以表格收尾时表格之后补一个空段。最小集与约定见 references/最小集.md，字段有改动先改那里。

默认写到临时位置（%TEMP%/loo0ng-to-docx/），由门禁的 --deliver 在通过后才一次性落进工作区（fail-closed）。

用哪个解释器归 agent（ADR-0018）：约束的是后端版本而不是哪一个 python，装法与四条约束见 SKILL.md 的
「跑得动转换器的环境」。版本对不上照常出件，只在回显里报出来。

用法：
  python md2docx.py <稿.md> --template <模板.docx> [--out <输出.docx>]

退出码：0 写出（stdout 第一行「已写出 <路径>」，第二行「出件环境：…」，其后是给人看的备注）；1 拒绝（最小集
之外的写法、表形与模板不合、模板打不开；原因在 stderr，什么都不写）；2 用法错误。
"""
import argparse
import copy
import datetime as _dt
import importlib.metadata
import os
import pathlib
import re
import secrets
import sys
import tempfile
from typing import Dict, List, Optional, Tuple

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TITLE_PREFIX = "# "
RIGHT_PREFIX = ":right: "
LEFT_PREFIX = ":left: "
MERGE_LEFT = "<"
MERGE_UP = "^"
TEMP_DIRNAME = "loo0ng-to-docx"
BACKEND = "python-docx"
MANIFEST = pathlib.Path(__file__).resolve().parent.parent / "requirements.txt"  # 随包分发的依赖清单
DEFAULT_PAGE_WIDTH = 11906  # A4，twips
DEFAULT_MARGIN = 1440
SEPARATOR_CELL = re.compile(r"^:?-+:?$")
# 最小集之外的写法，一律拒绝：确定性转换不猜模型的意思。
REJECTED = (
    (re.compile(r"^#(?! )"), "只有一级标题「# 」，没有多级标题"),
    (re.compile(r"^```"), "没有代码块"),
    (re.compile(r"^[-*+]\s"), "没有列表符号，编号写进正文"),
    (re.compile(r"^>"), "没有引用"),
    (re.compile(r"!\["), "没有图片"),
    (re.compile(r"\*\*|__"), "没有行内加粗，一切标注写进审查报告"),
    (re.compile(r"^:(?!right: |left: )\w+:"), "只有 :right: 与 :left: 两个前缀"),
)
# OOXML 里 pPr 与 tcPr 子元素的固定顺序，插元素时按它放，Word 才不报修复。
PPR_ORDER = ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
             "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
             "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange")
TCPR_ORDER = ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark")
STRIP_FROM_PPR = ("sectPr", "numPr", "pPrChange")
JC = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}


class Rejected(Exception):
    """转换拒绝：什么都不写。"""


# ---------------------------------------------------------------- Markdown 最小集

def parse_markdown(text: str) -> List[Tuple[str, object]]:
    """把最小集解析成块：('title'|'body'|'left'|'right', 文本) 或 ('table', 行列表)。一行一段，空行只是分隔。"""
    blocks: List[Tuple[str, object]] = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        for pattern, why in REJECTED:
            if pattern.search(line):
                raise Rejected("第 %d 行不在最小集里（%s）：%s" % (i + 1, why, line[:40]))
        if line.startswith("|"):
            rows: List[List[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(SEPARATOR_CELL.match(c) for c in cells):
                    rows.append(cells)
                i += 1
            if not rows:
                raise Rejected("第 %d 行起的表只有分隔行、没有内容" % (i + 1))
            width = len(rows[0])
            for r, row in enumerate(rows):
                if len(row) != width:
                    raise Rejected("表第 %d 行有 %d 格，首行有 %d 格；每行格数须相同" % (r + 1, len(row), width))
            blocks.append(("table", rows))
            continue
        if line.startswith(TITLE_PREFIX):
            blocks.append(("title", line[len(TITLE_PREFIX):].strip()))
        elif line.startswith(RIGHT_PREFIX):
            blocks.append(("right", line[len(RIGHT_PREFIX):].strip()))
        elif line.startswith(LEFT_PREFIX):
            blocks.append(("left", line[len(LEFT_PREFIX):].strip()))
        else:
            blocks.append(("body", line))
        i += 1
    if not blocks:
        raise Rejected("稿子是空的")
    return blocks


# ---------------------------------------------------------------- 模板里的格式原型

def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _jc(p) -> str:
    ppr = p.find(qn("w:pPr"))
    jc = ppr.find(qn("w:jc")) if ppr is not None else None
    return jc.get(qn("w:val")) if jc is not None else ""


def _has_first_line_indent(p) -> bool:
    ppr = p.find(qn("w:pPr"))
    ind = ppr.find(qn("w:ind")) if ppr is not None else None
    return ind is not None and (ind.get(qn("w:firstLine")) or ind.get(qn("w:firstLineChars"))) is not None


def _first_run_rpr(p):
    """段落里第一个有字的 run 的 rPr；没有则退到任一 run 的 rPr，再退到段落标记的 rPr。"""
    for r in p.iter(qn("w:r")):
        if _text(r):
            rpr = r.find(qn("w:rPr"))
            if rpr is not None:
                return rpr
    for r in p.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is not None:
            return rpr
    ppr = p.find(qn("w:pPr"))
    return ppr.find(qn("w:rPr")) if ppr is not None else None


def _clean_ppr(ppr):
    """段落格式原型：深拷贝并去掉不该带走的东西（分节符、编号、修订）。"""
    if ppr is None:
        return None
    ppr = copy.deepcopy(ppr)
    for tag in STRIP_FROM_PPR:
        el = ppr.find(qn("w:" + tag))
        if el is not None:
            ppr.remove(el)
    return ppr


def _derive_ppr(ppr, *, jc: Optional[str] = None, drop_indent: bool = False):
    """从原型派生一份 pPr：可改对齐、可去首行缩进。原型为空时按需新建。"""
    ppr = copy.deepcopy(ppr) if ppr is not None else OxmlElement("w:pPr")
    if jc is not None:
        ppr.jc_val = JC[jc]
    if drop_indent:
        ind = ppr.find(qn("w:ind"))
        if ind is not None:
            for attr in ("w:firstLine", "w:firstLineChars", "w:hanging", "w:hangingChars"):
                if ind.get(qn(attr)) is not None:
                    del ind.attrib[qn(attr)]
            if not ind.attrib:
                ppr.remove(ind)
    return ppr


class Prototype:
    def __init__(self, ppr, rpr):
        self.ppr = ppr
        self.rpr = rpr

    @classmethod
    def of(cls, p) -> "Prototype":
        rpr = _first_run_rpr(p)
        return cls(_clean_ppr(p.find(qn("w:pPr"))), copy.deepcopy(rpr) if rpr is not None else None)


class Template:
    """打开模板，取出段落格式原型（标题、正文、右对齐）、顶层表格与页面尺寸。"""

    def __init__(self, path: pathlib.Path):
        try:
            self.doc = docx.Document(str(path))
        except Exception as e:  # python-docx 打不开的都算模板坏了
            raise Rejected("模板打不开：%s（%s）" % (path, e))
        self.body = self.doc.element.body
        self.sect_pr = self.body.find(qn("w:sectPr"))
        if self.sect_pr is None:
            raise Rejected("模板没有 sectPr，不是能当载体的 DOCX：%s" % path)
        paragraphs = [p for p in self.body.findall(qn("w:p")) if _text(p).strip()]
        self.tables = list(self.body.findall(qn("w:tbl")))
        # 标题原型：第一个居中的有字段落（4-2、4-3 首段是说明段，恰好也居中同格式）；没有居中的就取第一段
        centered = [p for p in paragraphs if _jc(p) == "center"]
        title_src = (centered or paragraphs or [None])[0]
        self.title = Prototype.of(title_src) if title_src is not None else Prototype(None, None)
        plain = [p for p in paragraphs[1:] if _jc(p) not in ("center", "right")]
        indented = [p for p in plain if _has_first_line_indent(p)]
        body_src = (indented or plain or [None])[0]
        self.body_proto = Prototype.of(body_src) if body_src is not None else Prototype(None, None)
        rights = [p for p in paragraphs if _jc(p) == "right"]
        if rights:
            self.right = Prototype.of(rights[0])
        else:
            self.right = Prototype(_derive_ppr(self.body_proto.ppr, jc="right", drop_indent=True),
                                   copy.deepcopy(self.body_proto.rpr) if self.body_proto.rpr is not None else None)

    def page_metrics(self) -> Tuple[int, int, int]:
        pg = self.sect_pr.find(qn("w:pgSz"))
        mar = self.sect_pr.find(qn("w:pgMar"))
        width = int(pg.get(qn("w:w"))) if pg is not None and pg.get(qn("w:w")) else DEFAULT_PAGE_WIDTH
        left = int(mar.get(qn("w:left"))) if mar is not None and mar.get(qn("w:left")) else DEFAULT_MARGIN
        right = int(mar.get(qn("w:right"))) if mar is not None and mar.get(qn("w:right")) else DEFAULT_MARGIN
        return width, left, right

    def usable_width(self) -> int:
        width, left, right = self.page_metrics()
        return width - left - right


# ---------------------------------------------------------------- 写段与表

def _insert_ordered(parent, child, order):
    """按 OOXML 固定子元素顺序插入 child。"""
    tag = child.tag.split("}")[1]
    rank = order.index(tag)
    for i, existing in enumerate(parent):
        name = existing.tag.split("}")[1]
        if name in order and order.index(name) > rank:
            parent.insert(i, child)
            return
    parent.append(child)


def make_paragraph(ppr, rpr, text: str):
    p = OxmlElement("w:p")
    if ppr is not None:
        p.append(copy.deepcopy(ppr))
    if text:
        r = OxmlElement("w:r")
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        p.append(r)
    return p


def _grid_cells(tr) -> List[Tuple[object, int, int]]:
    """把模板一行展开成 (tc, 起始列, 跨列数)。"""
    out = []
    col = 0
    for tc in tr.findall(qn("w:tc")):
        tcpr = tc.find(qn("w:tcPr"))
        span_el = tcpr.find(qn("w:gridSpan")) if tcpr is not None else None
        span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
        out.append((tc, col, span))
        col += span
    return out


def _covering(cells, col):
    for tc, start, span in cells:
        if start <= col < start + span:
            return tc, start, span
    raise Rejected("模板行只有 %d 列，写不到第 %d 列" % (sum(s for _, _, s in cells), col + 1))


def _cell_paragraph_protos(tc, fallback_rpr) -> Tuple[object, object]:
    p = tc.find(qn("w:p"))
    if p is None:
        return None, copy.deepcopy(fallback_rpr) if fallback_rpr is not None else None
    ppr = _clean_ppr(p.find(qn("w:pPr")))
    rpr = _first_run_rpr(p)
    if rpr is None:
        rpr = fallback_rpr
    return ppr, (copy.deepcopy(rpr) if rpr is not None else None)


def _row_rpr(tr):
    for tc in tr.findall(qn("w:tc")):
        for p in tc.findall(qn("w:p")):
            rpr = _first_run_rpr(p)
            if rpr is not None:
                return rpr
    return None


def _tcw(tcpr):
    tcw = tcpr.find(qn("w:tcW")) if tcpr is not None else None
    if tcw is None or tcw.get(qn("w:w")) is None:
        return None, None
    return int(tcw.get(qn("w:w"))), tcw.get(qn("w:type")) or "dxa"


def _set_tcpr(tcpr, tag, attrs: Dict[str, str]):
    old = tcpr.find(qn("w:" + tag))
    if old is not None:
        tcpr.remove(old)
    el = OxmlElement("w:" + tag)
    for k, v in attrs.items():
        el.set(qn("w:" + k), v)
    _insert_ordered(tcpr, el, TCPR_ORDER)
    return el


def _group_cells(row: List[str], row_no: int) -> List[Tuple[int, int, str, bool]]:
    """把一行 Markdown 格按 `<`（并左）与 `^`（并上）分组：(起始列, 跨列, 文本, 是否并上)。"""
    groups: List[Tuple[int, int, str, bool]] = []
    for col, cell in enumerate(row):
        if cell == MERGE_LEFT:
            if not groups:
                raise Rejected("表第 %d 行第 1 格写了「<」，左边没有格可并" % row_no)
            start, span, text, up = groups[-1]
            groups[-1] = (start, span + 1, text, up)
        elif cell == MERGE_UP:
            groups.append((col, 1, "", True))
        else:
            groups.append((col, 1, cell, False))
    return groups


def _merge_up(tcpr, prev_cells: Dict[int, object], start: int, table_no: int, row_no: int) -> None:
    above = prev_cells.get(start)
    if above is None:
        raise Rejected("第 %d 张表第 %d 行第 %d 列写了「^」，上一行同一位置没有格可并" % (table_no, row_no, start + 1))
    above_pr = above.find(qn("w:tcPr"))
    if above_pr.find(qn("w:vMerge")) is None:
        _set_tcpr(above_pr, "vMerge", {"val": "restart"})
    _set_tcpr(tcpr, "vMerge", {})


def _build_rows(tbl, rows: List[List[str]], table_no: int, row_setup) -> None:
    """逐行逐格写表，合并由 `<`、`^` 约定决定。row_setup(r, row) 给出这一行的 trPr（可空）与
    make_cell(start, span) -> (tcPr, 段落 pPr, 字体 rPr)；模板表与自由表只差这两样。"""
    prev_cells: Dict[int, object] = {}
    for r, row in enumerate(rows):
        trpr, make_cell = row_setup(r, row)
        tr = OxmlElement("w:tr")
        if trpr is not None:
            tr.append(trpr)
        cur_cells: Dict[int, object] = {}
        for start, span, text, up in _group_cells(row, r + 1):
            tcpr, ppr, rpr = make_cell(start, span)
            if span > 1:
                _set_tcpr(tcpr, "gridSpan", {"val": str(span)})
            if up:
                _merge_up(tcpr, prev_cells, start, table_no, r + 1)
            tc = OxmlElement("w:tc")
            tc.append(tcpr)
            tc.append(make_paragraph(ppr, rpr, text))
            tr.append(tc)
            cur_cells[start] = tc
        tbl.append(tr)
        prev_cells = cur_cells


def build_template_table(tpl_tbl, rows: List[List[str]], table_no: int, fallback_rpr):
    """照抄模板表：tblPr 与 tblGrid 整个拷贝；第 i 行照模板第 i 行（超出的照最后一行）的行高与格格式。"""
    tbl = OxmlElement("w:tbl")
    tbl.append(copy.deepcopy(tpl_tbl.find(qn("w:tblPr"))))
    grid = tpl_tbl.find(qn("w:tblGrid"))
    tbl.append(copy.deepcopy(grid))
    ncols = len(grid.findall(qn("w:gridCol")))
    tpl_rows = tpl_tbl.findall(qn("w:tr"))
    if not tpl_rows:
        raise Rejected("模板第 %d 张表没有行" % table_no)

    def row_setup(r, row):
        if len(row) != ncols:
            raise Rejected("第 %d 张表第 %d 行有 %d 格，模板这张表是 %d 列；每格一列，合并用「<」「^」占位"
                           % (table_no, r + 1, len(row), ncols))
        base = tpl_rows[min(r, len(tpl_rows) - 1)]
        base_cells = _grid_cells(base)
        row_rpr = _row_rpr(base)
        if row_rpr is None:
            row_rpr = fallback_rpr
        trpr = base.find(qn("w:trPr"))

        def make_cell(start, span):
            src, src_start, src_span = _covering(base_cells, start)
            src_pr = src.find(qn("w:tcPr"))
            tcpr = copy.deepcopy(src_pr) if src_pr is not None else OxmlElement("w:tcPr")
            for tag in ("gridSpan", "vMerge", "hMerge"):
                el = tcpr.find(qn("w:" + tag))
                if el is not None:
                    tcpr.remove(el)
            if not (src_start == start and src_span == span):
                # 模板里没有正好这一格：宽度按所跨各列在模板里的宽度相加（同一行的宽度类型一致）
                total, wtype = 0, None
                for col in range(start, start + span):
                    c, _, csp = _covering(base_cells, col)
                    w, t = _tcw(c.find(qn("w:tcPr")))
                    if w is not None:
                        total += w // csp
                        wtype = wtype or t
                if wtype is not None:
                    _set_tcpr(tcpr, "tcW", {"w": str(total), "type": wtype})
            ppr, rpr = _cell_paragraph_protos(src, row_rpr)
            return tcpr, ppr, rpr

        return (copy.deepcopy(trpr) if trpr is not None else None), make_cell

    _build_rows(tbl, rows, table_no, row_setup)
    return tbl


def build_free_table(rows: List[List[str]], usable_width: int, table_no: int, rpr):
    """模板里没有对应的表：等宽铺满版心、细线边框、正文字体。"""
    ncols = len(rows[0])
    col_w = usable_width // ncols
    tbl = OxmlElement("w:tbl")
    tblpr = OxmlElement("w:tblPr")
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), "0")
    tblw.set(qn("w:type"), "auto")
    tblpr.append(tblw)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement("w:" + side)
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "auto")
        borders.append(b)
    tblpr.append(borders)
    tbl.append(tblpr)
    grid = OxmlElement("w:tblGrid")
    for _ in range(ncols):
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(col_w))
        grid.append(gc)
    tbl.append(grid)

    def make_cell(start, span):
        tcpr = OxmlElement("w:tcPr")
        _set_tcpr(tcpr, "tcW", {"w": str(col_w * span), "type": "dxa"})
        return tcpr, None, rpr

    _build_rows(tbl, rows, table_no, lambda r, row: (None, make_cell))
    return tbl


# ---------------------------------------------------------------- 转换

def convert(markdown: str, template_path: pathlib.Path, out_path: pathlib.Path) -> List[str]:
    """转换并写出；返回给人看的备注行（第几张表照抄了模板、哪些是自由表、有没有补空段）。"""
    blocks = parse_markdown(markdown)
    tpl = Template(template_path)
    body = tpl.body
    for child in list(body):
        if child is not tpl.sect_pr:
            body.remove(child)
    notes: List[str] = []
    body_ppr = tpl.body_proto.ppr
    left_ppr = _derive_ppr(body_ppr, drop_indent=True)
    table_no = 0
    last = None
    for kind, payload in blocks:
        if kind == "title":
            el = make_paragraph(tpl.title.ppr, tpl.title.rpr, payload)
        elif kind == "body":
            el = make_paragraph(body_ppr, tpl.body_proto.rpr, payload)
        elif kind == "left":
            el = make_paragraph(left_ppr, tpl.body_proto.rpr, payload)
        elif kind == "right":
            el = make_paragraph(tpl.right.ppr, tpl.right.rpr, payload)
        else:
            table_no += 1
            if table_no <= len(tpl.tables):
                el = build_template_table(tpl.tables[table_no - 1], payload, table_no, tpl.body_proto.rpr)
                notes.append("第 %d 张表照抄模板第 %d 张表的列宽、行高与格式" % (table_no, table_no))
            else:
                el = build_free_table(payload, tpl.usable_width(), table_no, tpl.body_proto.rpr)
                notes.append("第 %d 张表模板里没有对应的表，按等宽自由表写" % table_no)
        tpl.sect_pr.addprevious(el)
        last = el
    if last is not None and last.tag == qn("w:tbl"):
        tpl.sect_pr.addprevious(make_paragraph(left_ppr, None, ""))
        notes.append("正文以表格收尾，表后补了一个空段")
    clear_metadata(tpl.doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    part = out_path.with_name(out_path.name + ".part")
    tpl.doc.save(str(part))
    os.replace(part, out_path)
    return notes


def clear_metadata(doc) -> None:
    """通用裁定台账 #10：作者与最后修改者置空、修订号置 1、删上次打印时间、创建与修改时间置为本次生成时间。"""
    cp = doc.core_properties
    now = _dt.datetime.now(_dt.timezone.utc)
    cp.author = ""
    cp.last_modified_by = ""
    cp.revision = 1
    lp = cp._element.find("{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}lastPrinted")
    if lp is not None:
        cp._element.remove(lp)  # python-docx 不接受 None，须直接删元素
    cp.created = now
    cp.modified = now


def default_out_path(md_path: pathlib.Path) -> pathlib.Path:
    base = pathlib.Path(tempfile.gettempdir()) / TEMP_DIRNAME
    base.mkdir(parents=True, exist_ok=True)
    while True:
        candidate = base / ("%s-%s.docx" % (md_path.stem, secrets.token_hex(4)))
        if not candidate.exists():
            return candidate


# ---------------------------------------------------------------- 出件环境（ADR-0018）

def pinned_version() -> Optional[str]:
    """清单里钉的后端版本；读不到就 None。只用标准库：转换器的第三方 import 仍然只有 docx。"""
    try:
        text = MANIFEST.read_text(encoding="utf-8")
    except OSError:
        return None
    for raw in text.splitlines():
        line = raw.split("#", 1)[0].strip()
        name, sep, value = line.partition("==")
        if sep and name.strip().lower().replace("_", "-") == BACKEND:
            return value.strip()
    return None


def runtime_version() -> Optional[str]:
    """这次真正跑起来的 python-docx 版本，问运行时要，不照抄清单。"""
    try:
        return importlib.metadata.version(BACKEND)
    except importlib.metadata.PackageNotFoundError:
        return None


def environment_line() -> str:
    """回显里那一行出件环境：恒常写，模型原样抄进审查报告的「生成依据」段（ADR-0018）。

    版本对不上照常出件、退出码照旧是 0，只在这一行里明写它不是钉住的那个版本：缺能力不阻断，但不把没量过
    的说成量过了。
    """
    actual = runtime_version()
    pinned = pinned_version()
    if actual is None:
        backend = "%s 版本查不出来（这个解释器上没有它的包元数据）" % BACKEND
    elif pinned is None:
        backend = "%s %s（依赖清单读不到，无从比对）" % (BACKEND, actual)
    elif actual == pinned:
        backend = "%s %s（与依赖清单钉的一致）" % (BACKEND, actual)
    else:
        backend = "%s %s（依赖清单钉的是 %s，这不是钉住的那个版本）" % (BACKEND, actual, pinned)
    return "出件环境：解释器 %s（python %d.%d.%d）；%s" % (
        sys.executable, sys.version_info[0], sys.version_info[1], sys.version_info[2], backend)


# ---------------------------------------------------------------- CLI

def build_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(prog="md2docx.py", description="Markdown 最小集 → DOCX，以官方模板为版式载体。")
    ap.add_argument("markdown", help="模型写的稿子（.md，UTF-8）")
    ap.add_argument("--template", required=True, help="该节点的官方模板（.docx）")
    ap.add_argument("--out", help="输出路径；不给则写到 %%TEMP%%/loo0ng-to-docx/ 下并打印路径。已存在的文件不覆盖")
    return ap


def main(argv: Optional[List[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        md_path = pathlib.Path(args.markdown)
        if not md_path.is_file():
            raise Rejected("稿子不存在：%s" % md_path)
        template = pathlib.Path(args.template)
        if not template.is_file():
            raise Rejected("模板不存在：%s" % template)
        out = pathlib.Path(args.out) if args.out else default_out_path(md_path)
        if out.exists():
            raise Rejected("输出已存在，不覆盖：%s" % out)
        try:
            markdown = md_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            raise Rejected("稿子不是 UTF-8：%s" % md_path)
        notes = convert(markdown, template, out)
    except Rejected as e:
        sys.stderr.write("拒绝：%s\n" % e)
        return 1
    print("已写出 %s" % out)
    print(environment_line())
    for n in notes:
        print(n)
    return 0


if __name__ == "__main__":
    sys.exit(main())
