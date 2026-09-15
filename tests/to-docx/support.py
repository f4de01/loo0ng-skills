"""tests/to-docx 的公用件：定位仓库、模板目录与填模板 CLI，跑子进程，读 DOCX 里的 XML。

模板目录：官方模板原件住出厂预设图 skills/engineering/domain/assets/预设图/破产/模板/（ADR-0023，#29）。找不到就直接报错，不 skip。

只有一个 CLI（#13 起，#22 之后）：`fill.py` 打清单与施加差量（只依赖 python-docx）。版式门禁按 ADR-0024 整件退场，
这套测试不起 Word、任何机器上必须全绿、不许 skip。
造件一律「官方模板 + 一份差量」：模板就是载体，填出来的件与律师手里那份同一个骨架；要故障件再在它上面做 XML 手术。
`fill.py` 另按模块 import 一份（`填`），测试拿它算槽号、造差量、比几何与格式；那是库这一侧，不经命令行。
"""
import copy
import importlib.util
import json
import pathlib
import re
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET

REPO = pathlib.Path(__file__).resolve().parents[2]
SCRIPTS = REPO / "skills" / "productivity" / "to-docx" / "scripts"
FILL = SCRIPTS / "fill.py"
FIXTURES = pathlib.Path(__file__).resolve().parent / "fixtures"
PROBE_TEMPLATE = "1-2.关于管理人印章备案的报告.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
CP = "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}"
DC = "{http://purl.org/dc/elements/1.1/}"

_spec = importlib.util.spec_from_file_location("loo0ng_fill_for_tests", FILL)
填 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(填)

from docx import Document  # noqa: E402  测试侧随便 import；脚本侧的依赖边界由 test_isolation.py 守
from docx.oxml.ns import qn  # noqa: E402


def templates_dir() -> pathlib.Path:
    templates = REPO / "skills" / "engineering" / "domain" / "assets" / "预设图" / "破产" / "模板"
    if (templates / PROBE_TEMPLATE).is_file():
        return templates
    raise AssertionError("找不到官方模板目录 %s" % templates)


def template(prefix: str) -> pathlib.Path:
    """按文件名前缀（如 "1-2."）找一件官方模板。"""
    hits = sorted(p for p in templates_dir().glob("*.docx") if p.name.startswith(prefix))
    if len(hits) != 1:
        raise AssertionError("模板前缀 %s 命中 %d 件" % (prefix, len(hits)))
    return hits[0]


def all_templates():
    return sorted(templates_dir().glob("*.docx"))


class Run:
    def __init__(self, code, out, err):
        self.code, self.out, self.err = code, out, err

    def __repr__(self):
        return "Run(code=%r, out=%r, err=%r)" % (self.code, self.out, self.err)


def run(script, *args) -> Run:
    r = subprocess.run([sys.executable, str(script), *map(str, args)], capture_output=True, text=True,
                       encoding="utf-8", errors="replace")
    return Run(r.returncode, r.stdout, r.stderr)


# ---------------------------------------------------------------- 打清单与施加差量

def listing(docx: pathlib.Path, *extra) -> Run:
    return run(FILL, "list", docx, *extra)


def apply(src: pathlib.Path, ops, out: pathlib.Path, *extra) -> Run:
    """跑一次 fill.py apply：差量写成 JSON 落在 out 旁边，回显解析成 Run.changed（改动段号）。"""
    diff = out.with_suffix(".diff.json")
    diff.write_text(json.dumps(ops, ensure_ascii=False), encoding="utf-8")
    r = run(FILL, "apply", src, "--diff", diff, "--out", out, *extra)
    r.changed = changed_of(r.out)
    return r


def out_path(r: Run) -> pathlib.Path:
    """apply 回显第一行「已写出 <路径>」里的那个路径。"""
    first = r.out.splitlines()[0]
    assert first.startswith("已写出 "), "第一行不是产物路径：%r" % r.out
    return pathlib.Path(first.split(" ", 1)[1])


def changed_of(stdout: str):
    """回显里「改动段：2,3,13」那一行，解析成 [2, 3, 13]；「无」是空表。"""
    for line in stdout.splitlines():
        if line.startswith("改动段："):
            body = line.split("：", 1)[1].strip()
            return [] if body == "无" else [int(x) for x in body.split(",")]
    return None


def fill_all_ops(src: pathlib.Path, text: str = "某某"):
    """给一件 docx 造一份「每个槽都填上同一段字」的差量。"""
    doc = Document(str(src))
    ops = []
    for i, p in enumerate(填.paragraphs(doc)):
        for n, _ in enumerate(填.slots(p), 1):
            ops.append({"op": "fill", "at": "p%d#%d" % (i, n), "text": text})
    return ops


def slots_of(src: pathlib.Path):
    """[(段号, [(起, 止, 原文), …]), …]，只取有槽的段。"""
    doc = Document(str(src))
    out = []
    for i, p in enumerate(填.paragraphs(doc)):
        sl = 填.slots(p)
        if sl:
            out.append((i, sl))
    return out


def paragraph_count(src: pathlib.Path) -> int:
    return len(填.paragraphs(Document(str(src))))


# ---------------------------------------------------------------- 几何与格式快照（原地施加的验收）

def geometry_snapshot(src: pathlib.Path):
    """(表格几何, 每段每个字符的 run 格式)。

    表格几何 = 每张表的 tblPr / tblGrid、每行 trPr、每格 tcPr 的 XML；施加前后逐字节比。
    格式 = 每个字符所在 run 的 rPr（去掉 highlight），施加前后逐字符比：高亮本来就是这条路要加的东西。
    """
    doc = Document(str(src))
    geo = []
    for tbl in doc.element.body.iter(qn("w:tbl")):
        for tag in ("w:tblPr", "w:tblGrid"):
            el = tbl.find(qn(tag))
            geo.append(el.xml if el is not None else "")
        for tr in tbl.findall(qn("w:tr")):
            el = tr.find(qn("w:trPr"))
            geo.append(el.xml if el is not None else "")
            for tc in tr.findall(qn("w:tc")):
                el = tc.find(qn("w:tcPr"))
                geo.append(el.xml if el is not None else "")
    fmt = {}
    for idx, p in enumerate(填.paragraphs(doc)):
        chars = []
        for r, s, e in 填._runs(p):
            chars.extend([_rpr_key(r)] * (e - s))
        fmt[idx] = chars
    return geo, fmt


def _rpr_key(r) -> str:
    rpr = r.find(qn("w:rPr"))
    if rpr is None:
        return ""
    rpr = copy.deepcopy(rpr)
    el = rpr.find(qn("w:highlight"))
    if el is not None:
        rpr.remove(el)
    return rpr.xml


# ---------------------------------------------------------------- 读 DOCX

def read_xml(docx: pathlib.Path, member: str = "word/document.xml"):
    with zipfile.ZipFile(str(docx)) as z:
        return ET.fromstring(z.read(member))


def rewrite(src: pathlib.Path, dst: pathlib.Path, edits) -> pathlib.Path:
    """按成员名改写 DOCX 里的部件：edits 是 {成员名: fn(bytes) -> bytes}。构造故障件用。"""
    with zipfile.ZipFile(str(src)) as zin, zipfile.ZipFile(str(dst), "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in edits:
                data = edits[item.filename](data)
            zout.writestr(item, data)
    return dst


def text_of(el) -> str:
    return "".join(t.text or "" for t in el.iter(W + "t"))


def ending_with_a_table(src: pathlib.Path, dst: pathlib.Path):
    """造一件正文以「最后一张表 + 一段 + sectPr」收尾的 docx：表之后只留第一段，其余段落删掉。
    回 (路径, 那一段的段号)。守「表格直接接 sectPr」那条用。"""
    doc = Document(str(src))
    body = doc.element.body
    tbls = body.findall(qn("w:tbl"))
    assert tbls, "%s 里没有表" % src.name
    kept = None
    for el in list(tbls[-1].itersiblings()):
        if el.tag == qn("w:p") and kept is None:
            kept = el
        elif el.tag in (qn("w:p"), qn("w:tbl")):
            body.remove(el)
    assert kept is not None, "%s 最后一张表之后没有段落" % src.name
    doc.save(str(dst))
    return dst, 填.paragraphs(doc).index(kept)


def first_paragraph_after_last_table(src: pathlib.Path):
    """正文最后一张表之后第一段的段号，连同表后有几段。"""
    doc = Document(str(src))
    tbl = doc.element.body.findall(qn("w:tbl"))[-1]
    after = [el for el in tbl.itersiblings() if el.tag == qn("w:p")]
    return 填.paragraphs(doc).index(after[0]), len(after)


def metadata_leftovers(docx: pathlib.Path):
    """收尾该清掉的元数据里还剩着的：作者、最后修改者、上次打印时间。全清了就是空表。"""
    core = read_xml(docx, "docProps/core.xml")
    left = []
    for tag, 名 in ((DC + "creator", "作者"), (CP + "lastModifiedBy", "最后修改者")):
        el = core.find(tag)
        if el is not None and (el.text or "").strip():
            left.append("%s没清" % 名)
    if core.find(CP + "lastPrinted") is not None:
        left.append("上次打印时间没清")
    return left


def body_blocks(docx: pathlib.Path):
    body = read_xml(docx).find(W + "body")
    return [c for c in body if c.tag in (W + "p", W + "tbl")]


def table_signature(tbl):
    """表的形：列宽、每行行高、每格 (gridSpan, vMerge)。同形 = 签名相等。"""
    grid = [g.get(W + "w") for g in tbl.find(W + "tblGrid").findall(W + "gridCol")]
    rows = []
    for tr in tbl.findall(W + "tr"):
        trpr = tr.find(W + "trPr")
        h = trpr.find(W + "trHeight") if trpr is not None else None
        height = (h.get(W + "val"), h.get(W + "hRule")) if h is not None else None
        cells = []
        for tc in tr.findall(W + "tc"):
            tcpr = tc.find(W + "tcPr")
            span = tcpr.find(W + "gridSpan") if tcpr is not None else None
            vm = tcpr.find(W + "vMerge") if tcpr is not None else None
            cells.append((span.get(W + "val") if span is not None else "1",
                          (vm.get(W + "val") or "continue") if vm is not None else None))
        rows.append((height, tuple(cells)))
    return (tuple(grid), tuple(rows))


def strip_ws(s: str) -> str:
    return re.sub(r"\s+", "", s)
